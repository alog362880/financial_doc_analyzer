"""Builds CSV and Excel audit reports from stored analyses."""

import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Document

DOCUMENT_HEADERS = [
    "Document ID",
    "Filename",
    "Status",
    "Pages",
    "Risk Score",
    "Risk Level",
    "Flagged Clauses",
    "Uploaded At",
]
CLAUSE_HEADERS = [
    "Document ID",
    "Filename",
    "Risk Score",
    "Risk Level",
    "Severity",
    "Category",
    "Page",
    "Clause Text",
    "Mitigation Suggestion",
    "Source Verified",
]

_FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True, slots=True)
class ReportTables:
    documents: list[list[Any]]
    clauses: list[list[Any]]


def _safe_cell(value: Any) -> Any:
    """Neutralise spreadsheet formula injection for text that came from uploaded documents or the LLM."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def build_report_tables(documents: Sequence[Document]) -> ReportTables:
    """Flatten ORM objects into plain rows so rendering can run off the event loop."""
    document_rows: list[list[Any]] = []
    clause_rows: list[list[Any]] = []

    for document in documents:
        analysis = document.analysis
        clauses = analysis.flagged_clauses if analysis else []
        document_rows.append(
            [
                document.id,
                _safe_cell(document.original_filename),
                document.status,
                document.page_count,
                analysis.risk_score if analysis else None,
                analysis.risk_level if analysis else None,
                len(clauses),
                document.created_at.isoformat(),
            ]
        )
        for clause in clauses:
            clause_rows.append(
                [
                    document.id,
                    _safe_cell(document.original_filename),
                    analysis.risk_score if analysis else None,
                    analysis.risk_level if analysis else None,
                    clause.get("severity"),
                    clause.get("category"),
                    clause.get("page_number"),
                    _safe_cell(clause.get("clause_text")),
                    _safe_cell(clause.get("mitigation_suggestion")),
                    "Yes" if clause.get("source_verified") else "No",
                ]
            )

    return ReportTables(documents=document_rows, clauses=clause_rows)


def render_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    # UTF-8 with BOM so Excel detects encoding correctly.
    return buffer.getvalue().encode("utf-8-sig")


def _write_sheet(sheet: Worksheet, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def render_xlsx(tables: ReportTables) -> bytes:
    workbook = Workbook()
    documents_sheet = workbook.active
    if documents_sheet is None:  # pragma: no cover - a new Workbook always has an active sheet
        raise RuntimeError("Workbook has no active sheet")
    documents_sheet.title = "Documents"
    _write_sheet(documents_sheet, DOCUMENT_HEADERS, tables.documents, [38, 40, 14, 10, 12, 14, 16, 26])

    clauses_sheet = workbook.create_sheet("Flagged Clauses")
    _write_sheet(clauses_sheet, CLAUSE_HEADERS, tables.clauses, [38, 36, 12, 12, 12, 24, 8, 80, 60, 14])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
